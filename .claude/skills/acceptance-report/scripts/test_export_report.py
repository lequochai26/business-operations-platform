import unicodedata

import openpyxl
import pytest
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image

from export_report import build_pdf, build_xlsx, collect_runs, latest_by_case, main, parse_results

RESULTS = """# Kết quả chạy test 2026-11-10-01
- Môi trường: máy dev, http://localhost:3000
- Commit: abc1234
- Người chạy: Claude + Hải
- Ngày: 2026-11-10

## P0-GRP-012 Giới hạn comment/ngày = 51
- Kết quả: không đạt
- Thực tế: lưu được giá trị 51
- Evidence: P0-GRP-012_01_input.png, P0-GRP-012_02_saved.png
- Ghi chú:
- Vấn đề: P0-ISS-003
"""


def write_run(phase_dir, run_id, body, images=()):
    run_dir = phase_dir / "test" / "runs" / run_id
    (run_dir / "evidence").mkdir(parents=True)
    (run_dir / "results.md").write_text(body, encoding="utf-8")
    for name, size in images:
        Image.new("RGB", size, "#8FA284").save(run_dir / "evidence" / name)
    return run_dir


def one_case(run_id, case_id, result, evidence=""):
    return (
        f"# Kết quả chạy test {run_id}\n- Môi trường: dev\n- Commit: abc1234\n- Người chạy: Claude\n- Ngày: 2026-11-11\n\n"
        f"## {case_id} Tiêu đề\n- Kết quả: {result}\n- Thực tế:\n- Evidence: {evidence}\n- Ghi chú:\n- Vấn đề:\n"
    )


@pytest.fixture
def phase_dir(tmp_path):
    d = tmp_path / "phase-0"
    (d / "acceptance").mkdir(parents=True)
    (d / "acceptance" / "report.md").write_text("# Báo cáo nghiệm thu – Giai đoạn 0\n\n## Tổng quan\n\nMẫu mới đón 20/10.\n", encoding="utf-8")
    (d / "issues.md").write_text("## P0-ISS-003 Lỗi\n- Trạng thái: mở\n\n## P0-ISS-004 Khác\n- Trạng thái: đã xong\n", encoding="utf-8")
    return d


def test_parse_results_reads_header_and_cases(tmp_path):
    run_dir = write_run(tmp_path, "2026-11-10-01", RESULTS)
    run = parse_results(run_dir / "results.md")
    assert run.run_id == "2026-11-10-01"
    assert run.commit == "abc1234"
    assert run.tester == "Claude + Hải"
    case = run.cases[0]
    assert case.case_id == "P0-GRP-012"
    assert case.title == "Giới hạn comment/ngày = 51"
    assert case.result == "không đạt"
    assert case.evidence == ["P0-GRP-012_01_input.png", "P0-GRP-012_02_saved.png"]
    assert case.issue == "P0-ISS-003"
    assert case.note == ""


def test_parse_results_handles_crlf_and_bom(tmp_path):
    run_dir = write_run(tmp_path, "2026-11-10-01", "﻿" + RESULTS.replace("\n", "\r\n"))
    run = parse_results(run_dir / "results.md")
    assert run.run_id == "2026-11-10-01"
    assert run.cases[0].issue == "P0-ISS-003"


def test_parse_results_normalizes_result(tmp_path):
    nfd = unicodedata.normalize("NFD", "đạt")
    run_dir = write_run(tmp_path, "a", one_case("a", "P0-GRP-001", "  Đạt ") + one_case("a", "P0-GRP-002", nfd).split("\n\n", 1)[1])
    results = [c.result for c in parse_results(run_dir / "results.md").cases]
    assert results == ["đạt", "đạt"]


def test_parse_results_rejects_unknown_result(tmp_path):
    run_dir = write_run(tmp_path, "a", one_case("a", "P0-GRP-001", "ok"))
    with pytest.raises(ValueError) as exc:
        parse_results(run_dir / "results.md")
    assert "results.md" in str(exc.value)
    assert "dòng 8" in str(exc.value)


def test_latest_by_case_prefers_later_run(phase_dir):
    write_run(phase_dir, "2026-11-10-01", one_case("2026-11-10-01", "P0-GRP-012", "không đạt"))
    write_run(phase_dir, "2026-11-11-01", one_case("2026-11-11-01", "P0-GRP-012", "đạt"))
    latest = latest_by_case(collect_runs(phase_dir))
    run, case = latest["P0-GRP-012"]
    assert run.run_id == "2026-11-11-01"
    assert case.result == "đạt"


def test_build_xlsx_creates_sheets(phase_dir):
    write_run(phase_dir, "2026-11-10-01", one_case("2026-11-10-01", "P0-GRP-012", "không đạt", "big.png"), [("big.png", (1200, 600))])
    out = phase_dir / "acceptance" / "report.xlsx"
    warnings = build_xlsx(phase_dir, collect_runs(phase_dir), out)
    assert warnings == []
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Tổng hợp", "Danh sách case", "P0-GRP-012"]
    images = wb["P0-GRP-012"]._images
    assert len(images) == 1
    assert images[0].anchor.ext.width == pixels_to_EMU(900)
    assert images[0].anchor.ext.height == pixels_to_EMU(450)


def test_build_xlsx_summary_counts_open_issues(phase_dir):
    write_run(phase_dir, "2026-11-10-01", one_case("2026-11-10-01", "P0-GRP-012", "đạt"))
    out = phase_dir / "acceptance" / "report.xlsx"
    build_xlsx(phase_dir, collect_runs(phase_dir), out)
    ws = openpyxl.load_workbook(out)["Tổng hợp"]
    rows = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r[0]}
    assert rows["đạt"] == 1
    assert rows["Vấn đề còn mở"] == 1


def test_build_xlsx_missing_evidence_warns(phase_dir):
    write_run(phase_dir, "2026-11-10-01", one_case("2026-11-10-01", "P0-GRP-012", "không đạt", "gone.png"))
    out = phase_dir / "acceptance" / "report.xlsx"
    warnings = build_xlsx(phase_dir, collect_runs(phase_dir), out)
    assert any("gone.png" in w for w in warnings)
    ws = openpyxl.load_workbook(out)["P0-GRP-012"]
    values = [c for row in ws.iter_rows(values_only=True) for c in row if c]
    assert "Thiếu file evidence: gone.png" in values


def test_build_pdf_writes_pdf(phase_dir):
    write_run(phase_dir, "2026-11-10-01", one_case("2026-11-10-01", "P0-GRP-012", "đạt", "a.png"), [("a.png", (400, 300))])
    out = phase_dir / "acceptance" / "report.pdf"
    assert build_pdf(phase_dir, collect_runs(phase_dir), out) == []
    assert out.read_bytes()[:4] == b"%PDF"


def test_build_pdf_missing_evidence_warns(phase_dir):
    write_run(phase_dir, "2026-11-10-01", one_case("2026-11-10-01", "P0-GRP-012", "đạt", "gone.png"))
    out = phase_dir / "acceptance" / "report.pdf"
    warnings = build_pdf(phase_dir, collect_runs(phase_dir), out)
    assert any("gone.png" in w for w in warnings)
    assert out.exists()


def test_main_reports_locked_output(phase_dir, monkeypatch, capsys):
    write_run(phase_dir, "2026-11-10-01", one_case("2026-11-10-01", "P0-GRP-012", "đạt"))

    def locked(self, filename):
        raise PermissionError(13, "Permission denied", str(filename))

    monkeypatch.setattr(openpyxl.Workbook, "save", locked)
    assert main([str(phase_dir)]) == 1
    err = capsys.readouterr().err
    assert "Đóng file" in err
    assert "report.xlsx" in err
