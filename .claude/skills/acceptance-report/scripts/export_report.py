"""Xuất báo cáo nghiệm thu của một giai đoạn ra Excel và PDF.

Chạy: py .claude/skills/acceptance-report/scripts/export_report.py docs/phases/phase-0
Đọc: test/runs/*/results.md, issues.md, acceptance/report.md. Ghi: acceptance/report.xlsx, acceptance/report.pdf.
"""

import math
import re
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import markdown
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright

RESULTS = ("đạt", "không đạt", "bị chặn", "bỏ qua")
OPEN_ISSUE_STATES = ("mở", "đang xử lý")
MAX_IMAGE_WIDTH = 900
ROW_HEIGHT_PX = 20
TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates"


@dataclass
class CaseResult:
    case_id: str
    title: str
    result: str
    actual: str = ""
    evidence: list[str] = field(default_factory=list)
    note: str = ""
    issue: str | None = None


@dataclass
class Run:
    run_id: str
    environment: str
    commit: str
    tester: str
    date: str
    cases: list[CaseResult]
    dir: Path


def _norm(text: str) -> str:
    return unicodedata.normalize("NFC", text).strip()


def _read(path: Path) -> str:
    return path.read_text(encoding="utf-8").lstrip("﻿").replace("\r\n", "\n")


def _field(line: str) -> tuple[str, str] | None:
    match = re.match(r"^- ([^:]+):(.*)$", line)
    return (_norm(match.group(1)), _norm(match.group(2))) if match else None


def parse_results(path: Path) -> Run:
    header: dict[str, str] = {}
    run_id = ""
    cases: list[CaseResult] = []
    for lineno, raw in enumerate(_read(path).split("\n"), start=1):
        line = _norm(raw)
        if line.startswith("# Kết quả chạy test"):
            run_id = _norm(line.removeprefix("# Kết quả chạy test"))
        elif line.startswith("## "):
            case_id, _, title = line[3:].partition(" ")
            cases.append(CaseResult(case_id=case_id, title=title.strip(), result=""))
        elif (kv := _field(line)) is not None:
            key, value = kv
            if not cases:
                header[key] = value
                continue
            case = cases[-1]
            if key == "Kết quả":
                result = value.lower()
                if result not in RESULTS:
                    raise ValueError(f"{path}, dòng {lineno}: Kết quả '{value}' không hợp lệ (chỉ dùng: {', '.join(RESULTS)})")
                case.result = result
            elif key == "Thực tế":
                case.actual = value
            elif key == "Evidence":
                case.evidence = [e.strip() for e in value.split(",") if e.strip()]
            elif key == "Ghi chú":
                case.note = value
            elif key == "Vấn đề":
                case.issue = value or None
    for case in cases:
        if not case.result:
            raise ValueError(f"{path}: case {case.case_id} thiếu dòng '- Kết quả:'")
    return Run(
        run_id=run_id or path.parent.name,
        environment=header.get("Môi trường", ""),
        commit=header.get("Commit", ""),
        tester=header.get("Người chạy", ""),
        date=header.get("Ngày", ""),
        cases=cases,
        dir=path.parent,
    )


def collect_runs(phase_dir: Path) -> list[Run]:
    runs = [parse_results(p) for p in (phase_dir / "test" / "runs").glob("*/results.md")]
    return sorted(runs, key=lambda r: r.run_id)


def latest_by_case(runs: list[Run]) -> dict[str, tuple[Run, CaseResult]]:
    latest: dict[str, tuple[Run, CaseResult]] = {}
    for run in sorted(runs, key=lambda r: r.run_id):
        for case in run.cases:
            latest[case.case_id] = (run, case)
    return dict(sorted(latest.items()))


def _open_issues(phase_dir: Path) -> list[str]:
    path = phase_dir / "issues.md"
    if not path.exists():
        return []
    open_ids, current = [], None
    for line in _read(path).split("\n"):
        line = _norm(line)
        if line.startswith("## "):
            current = line[3:].split(" ", 1)[0]
        elif (kv := _field(line)) and kv[0] == "Trạng thái" and kv[1].lower() in OPEN_ISSUE_STATES and current:
            open_ids.append(current)
    return open_ids


def _evidence_paths(run: Run, case: CaseResult) -> list[tuple[str, Path | None]]:
    items = []
    for name in case.evidence:
        path = run.dir / "evidence" / name
        items.append((name, path if path.exists() else None))
    return items


def _missing_warning(case: CaseResult, name: str) -> str:
    return f"{case.case_id}: thiếu file evidence {name}"


def build_xlsx(phase_dir: Path, runs: list[Run], out: Path) -> list[str]:
    latest = latest_by_case(runs)
    warnings: list[str] = []
    wb = Workbook()
    bold = Font(bold=True)

    summary = wb.active
    summary.title = "Tổng hợp"
    summary.append(["Giai đoạn", phase_dir.name])
    summary.append(["Số lần chạy", len(runs)])
    summary.append(["Số case", len(latest)])
    for result in RESULTS:
        summary.append([result, sum(1 for _, c in latest.values() if c.result == result)])
    open_issues = _open_issues(phase_dir)
    summary.append(["Vấn đề còn mở", len(open_issues)])
    summary.append(["Mã vấn đề còn mở", ", ".join(open_issues)])
    summary.append([])
    summary.append(["Theo tính năng", *RESULTS])
    features = sorted({cid.rsplit("-", 1)[0] for cid in latest})
    for feature in features:
        cases = [c for cid, (_, c) in latest.items() if cid.rsplit("-", 1)[0] == feature]
        summary.append([feature, *(sum(1 for c in cases if c.result == r) for r in RESULTS)])
    summary.column_dimensions["A"].width = 24

    listing = wb.create_sheet("Danh sách case")
    listing.append(["Mã", "Tiêu đề", "Kết quả", "Lần chạy", "Vấn đề", "Ghi chú"])
    for cell in listing[1]:
        cell.font = bold
    for run, case in latest.values():
        listing.append([case.case_id, case.title, case.result, run.run_id, case.issue or "", case.note])
    listing.column_dimensions["A"].width = 16
    listing.column_dimensions["B"].width = 50

    for run, case in latest.values():
        ws = wb.create_sheet(case.case_id[:31])
        ws.append([f"{case.case_id} {case.title}"])
        ws["A1"].font = bold
        ws.append(["Kết quả", case.result])
        ws.append(["Lần chạy", run.run_id])
        ws.append(["Thực tế", case.actual])
        row = 6
        for name, path in _evidence_paths(run, case):
            ws.cell(row=row, column=1, value=name if path else f"Thiếu file evidence: {name}")
            row += 1
            if path is None:
                warnings.append(_missing_warning(case, name))
                continue
            img = XlImage(str(path))
            if img.width > MAX_IMAGE_WIDTH:
                img.height = round(img.height * MAX_IMAGE_WIDTH / img.width)
                img.width = MAX_IMAGE_WIDTH
            ws.add_image(img, f"A{row}")
            row += math.ceil(img.height / ROW_HEIGHT_PX) + 2

    wb.save(out)
    return warnings


def build_pdf(phase_dir: Path, runs: list[Run], out: Path) -> list[str]:
    warnings: list[str] = []
    report_md = phase_dir / "acceptance" / "report.md"
    report_html = markdown.markdown(_read(report_md), extensions=["tables"]) if report_md.exists() else ""
    cases = []
    for run, case in latest_by_case(runs).values():
        evidence = []
        for name, path in _evidence_paths(run, case):
            if path is None:
                warnings.append(_missing_warning(case, name))
            evidence.append({"name": name, "uri": path.resolve().as_uri() if path else None})
        cases.append({**case.__dict__, "issue": case.issue or "", "run_id": run.run_id, "evidence": evidence})

    env = Environment(loader=FileSystemLoader(TEMPLATE_DIR), autoescape=select_autoescape(["html"]))
    html = env.get_template("report.html").render(title=f"Báo cáo nghiệm thu – {phase_dir.name}", report_html=report_html, cases=cases)

    with tempfile.TemporaryDirectory() as tmp:
        html_path = Path(tmp) / "report.html"
        html_path.write_text(html, encoding="utf-8")
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.goto(html_path.as_uri())
            page.wait_for_load_state("networkidle")
            page.pdf(path=str(out), format="A4", print_background=True, margin={"top": "15mm", "bottom": "15mm", "left": "12mm", "right": "12mm"})
            browser.close()
    return warnings


def main(argv: list[str]) -> int:
    if len(argv) != 1:
        print("Cách dùng: py .claude/skills/acceptance-report/scripts/export_report.py docs/phases/phase-N", file=sys.stderr)
        return 2
    phase_dir = Path(argv[0])
    out_dir = phase_dir / "acceptance"
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        runs = collect_runs(phase_dir)
        warnings = build_xlsx(phase_dir, runs, out_dir / "report.xlsx")
        warnings += build_pdf(phase_dir, runs, out_dir / "report.pdf")
    except PermissionError as exc:
        print(f"Không ghi được {exc.filename}. Đóng file {exc.filename} rồi chạy lại.", file=sys.stderr)
        return 1
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    for warning in dict.fromkeys(warnings):
        print(f"Cảnh báo: {warning}")
    print(f"Đã ghi {out_dir / 'report.xlsx'} và {out_dir / 'report.pdf'}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
